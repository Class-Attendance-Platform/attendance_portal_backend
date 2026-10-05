"""Report exports (API v2 section 8): Dhaka time, join dates, split PDF grid, header, every format."""
import base64
import csv
import io
import re
import zlib
from datetime import date, datetime, timedelta, timezone as dt_timezone
from unittest import mock

from django.test import TestCase

from apps.academic.models import StudentClassroom
from apps.attendance.models import AttendanceLog, AttendanceSession
from apps.attendance.tests.helpers import (
    client_for, make_admin, make_course_info, make_semester, make_student, make_teacher,
)
from apps.reports.generators import build_report, date_chunks, dates_per_page

D1, D2, D3 = date(2026, 9, 1), date(2026, 9, 8), date(2026, 9, 15)
# 20:30 UTC is 02:30 the next morning in Dhaka (UTC+6)
FIXED_NOW = datetime(2026, 10, 5, 20, 30, tzinfo=dt_timezone.utc)


def pdf_text(content: bytes) -> str:
    """The text drawn in a reportlab PDF (its page streams are ASCII85 + Flate encoded)."""
    parts = []
    for raw in re.findall(rb'stream\r?\n(.*?)endstream', content, re.S):
        try:
            parts.append(zlib.decompress(base64.a85decode(raw.strip(), adobe=True)).decode('latin-1'))
        except ValueError:
            continue
    return '\n'.join(parts)


def pdf_pages(content: bytes) -> int:
    return len(re.findall(rb'/Type /Page\b(?!s)', content))


def log(ci, student, day, status, session=None):
    return AttendanceLog.objects.create(
        course_info=ci, student=student, date=day, status=status, session=session,
        method=AttendanceLog.Method.QR if status == 'PRESENT' else '',
    )


class ExportTestCase(TestCase):
    """
    Level 3 Term I 2025-26, CSE301, classes on D1, D2, D3:
    - Ayesha (2302001) from the start: present, absent, present (two sessions on D3) -> 2/3
    - Tanvir (2302002) joined on D2 (late): blank, present, absent -> 1/2
    - Nusrat (2302003) left on D3 (not on an active semester's class list)
    - Rafi (2302004) soft-deleted account (never listed)
    """

    def setUp(self):
        self.teacher = make_teacher('teacher@example.com', first_name='Karim', last_name='Uddin')
        self.ayesha = make_student('a@example.com', 2302001, first_name='Ayesha', last_name='Rahman')
        self.tanvir = make_student('t@example.com', 2302002, first_name='Tanvir', last_name='Hasan')
        self.nusrat = make_student('n@example.com', 2302003, first_name='Nusrat', last_name='Jahan')
        self.rafi = make_student('r@example.com', 2302004, first_name='Rafi', last_name='Islam')
        self.semester = make_semester(students=[self.ayesha, self.tanvir, self.nusrat, self.rafi])
        self.ci = make_course_info(self.teacher, 'CSE301', semester=self.semester)
        self.ci.course.title = 'Software Engineering'
        self.ci.course.save()
        StudentClassroom.objects.filter(student=self.tanvir).update(joined_at=D2)
        StudentClassroom.objects.filter(student=self.nusrat).update(left_at=D3)
        self.rafi.user.deleted = True
        self.rafi.user.save()

        log(self.ci, self.ayesha, D1, 'PRESENT')
        log(self.ci, self.nusrat, D1, 'PRESENT')
        log(self.ci, self.rafi, D1, 'PRESENT')
        log(self.ci, self.ayesha, D2, 'ABSENT')
        log(self.ci, self.tanvir, D2, 'PRESENT')
        log(self.ci, self.nusrat, D2, 'ABSENT')
        # D3: two sessions; a present log on either one counts (one date = one class)
        first = AttendanceSession.objects.create(course_info=self.ci, date=D3, mode='QR_ONLINE', is_active=False)
        second = AttendanceSession.objects.create(course_info=self.ci, date=D3, mode='FACE', is_active=False)
        log(self.ci, self.ayesha, D3, 'ABSENT', session=first)
        log(self.ci, self.ayesha, D3, 'PRESENT', session=second)
        log(self.ci, self.tanvir, D3, 'ABSENT', session=first)

        self.client = client_for(self.teacher.user)

    def export(self, query, client=None):
        return (client or self.client).get(f'/api/reports/course-info/{self.ci.id}/export/?{query}')


class ReportDataTests(ExportTestCase):
    def test_rows_follow_the_class_list_and_join_dates(self):
        report = build_report(self.ci)
        self.assertEqual(report.dates, [D1, D2, D3])
        rows = {r.student_id: r for r in report.rows}
        self.assertEqual(sorted(rows), [2302001, 2302002])  # left and deleted accounts are not listed
        ayesha, tanvir = rows[2302001], rows[2302002]
        self.assertEqual(ayesha.cells, ['PRESENT', 'ABSENT', 'PRESENT'])
        self.assertEqual((ayesha.held, ayesha.attended, ayesha.percent, ayesha.below_min), (3, 2, 66.7, True))
        self.assertEqual(tanvir.cells, ['', 'PRESENT', 'ABSENT'])  # not in the class on D1: blank
        self.assertEqual((tanvir.held, tanvir.attended, tanvir.percent), (2, 1, 50.0))

    def test_header_values(self):
        report = build_report(self.ci)
        self.assertEqual(report.title, 'CSE301 — Software Engineering')
        self.assertEqual(report.info_lines()[:3], [
            'Semester: Level 3 · Term I · 2025-26',
            'Department: Computer Science and Engineering (CSE)',
            'Teacher: Karim Uddin',
        ])
        self.assertEqual(report.scope_text, 'Classes: 3 (01 Sep 2026 to 15 Sep 2026)')

    def test_generated_time_is_dhaka_time(self):
        with mock.patch('django.utils.timezone.now', return_value=FIXED_NOW):
            report = build_report(self.ci)
        self.assertEqual(report.generated_text, '06 Oct 2026, 02:30 AM (Dhaka time)')

    def test_finished_semester_lists_former_members_with_blanks_after_they_left(self):
        self.semester.is_active = False
        self.semester.save()
        report = build_report(self.ci)
        rows = {r.student_id: r for r in report.rows}
        self.assertEqual(sorted(rows), [2302001, 2302002, 2302003])
        nusrat = rows[2302003]
        self.assertEqual(nusrat.cells, ['PRESENT', 'ABSENT', ''])
        self.assertEqual((nusrat.held, nusrat.attended, nusrat.percent), (2, 1, 50.0))
        self.assertEqual(report.semester_text, 'Level 3 · Term I · 2025-26 (finished)')

    def test_one_date(self):
        report = build_report(self.ci, D1)
        self.assertEqual(report.dates, [D1])
        self.assertEqual({r.student_id: r.cells for r in report.rows}, {2302001: ['PRESENT'], 2302002: ['']})

    def test_a_date_without_a_class_is_blank_for_everyone(self):
        report = build_report(self.ci, date(2026, 9, 2))
        self.assertFalse(report.class_on_date)
        self.assertEqual([r.cells for r in report.rows], [[''], ['']])
        self.assertEqual(report.scope_text, 'Date: 02 Sep 2026 (no class on this date)')

    def test_no_teacher(self):
        self.ci.teacher = None
        self.ci.save()
        self.assertEqual(build_report(self.ci).teacher, 'Not assigned')

    def test_date_chunks(self):
        self.assertEqual(date_chunks([], 5), [[]])
        self.assertEqual([len(c) for c in date_chunks(list(range(12)), 5)], [5, 5, 2])
        self.assertGreaterEqual(dates_per_page(), 10)  # landscape A4 fits a useful number of dates


class ExportEndpointTests(ExportTestCase):
    def test_format_parameter_names_the_file_type(self):
        for fmt, content_type in (
            ('csv', 'text/csv'),
            ('pdf', 'application/pdf'),
            ('xlsx', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            ('docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
        ):
            res = self.export(f'format={fmt}')
            self.assertEqual(res.status_code, 200, fmt)
            self.assertEqual(res['Content-Type'], content_type)
            self.assertIn(f'CSE301_attendance_full.{fmt}', res['Content-Disposition'])

    def test_old_export_format_parameter_still_works(self):
        res = self.export('export_format=pdf')
        self.assertEqual(res['Content-Type'], 'application/pdf')
        self.assertEqual(self.export('')['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')  # default xlsx

    def test_bad_format_or_date(self):
        res = self.export('format=png')
        self.assertEqual((res.status_code, res.data['code']), (400, 'invalid_format'))
        self.assertIn('pdf, xlsx, csv or docx', res.data['message'])
        res = self.export('format=csv&date=05-10-2026')
        self.assertEqual((res.status_code, res.data['code']), (400, 'invalid_date'))

    def test_access(self):
        admin = client_for(make_admin('admin@example.com'))
        self.assertEqual(self.export('format=csv', admin).status_code, 200)
        other = client_for(make_teacher('other@example.com').user)
        res = self.export('format=pdf', other)
        self.assertEqual((res.status_code, res.data['message']), (403, 'You do not teach this course.'))
        self.assertEqual(self.export('format=csv', client_for(self.ayesha.user)).status_code, 403)
        from rest_framework.test import APIClient
        self.assertEqual(self.export('format=pdf', APIClient()).status_code, 401)

    def test_deleted_course_or_semester_is_not_found(self):
        self.semester.deleted = True
        self.semester.save()
        res = self.export('format=pdf')
        self.assertEqual((res.status_code, res.data['code']), (404, 'not_found'))
        self.semester.deleted = False
        self.semester.save()
        self.ci.course.deleted = True
        self.ci.course.save()
        self.assertEqual(self.export('format=csv').status_code, 404)

    def test_csv(self):
        rows = list(csv.reader(io.StringIO(self.export('format=csv').content.decode())))
        self.assertEqual(rows[0], ['Student ID', 'Name', 'Email', '2026-09-01', '2026-09-08', '2026-09-15',
                                   'Classes held', 'Present', 'Percentage'])
        self.assertEqual(rows[1], ['2302001', 'Ayesha Rahman', 'a@example.com',
                                   'PRESENT', 'ABSENT', 'PRESENT', '3', '2', '66.7%'])
        self.assertEqual(rows[2], ['2302002', 'Tanvir Hasan', 't@example.com',
                                   '', 'PRESENT', 'ABSENT', '2', '1', '50.0%'])
        self.assertEqual(len(rows), 3)

    def test_csv_one_date(self):
        rows = list(csv.reader(io.StringIO(self.export('format=csv&date=2026-09-01').content.decode())))
        self.assertEqual(rows, [['Student ID', 'Name', 'Email', 'Status'],
                                ['2302001', 'Ayesha Rahman', 'a@example.com', 'PRESENT'],
                                ['2302002', 'Tanvir Hasan', 't@example.com', '']])

    def test_xlsx(self):
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(self.export('format=xlsx').content)).active
        values = list(ws.iter_rows(values_only=True))
        header_text = [row[0] for row in values[:7]]
        self.assertEqual(header_text[0], 'CSE301 — Software Engineering')
        self.assertIn('Semester: Level 3 · Term I · 2025-26', header_text)
        self.assertIn('Department: Computer Science and Engineering (CSE)', header_text)
        self.assertIn('Teacher: Karim Uddin', header_text)
        header_index = next(i for i, row in enumerate(values) if row[0] == '#')
        self.assertEqual(values[header_index][:8], ('#', 'Student ID', 'Name', 'Email',
                                                    '01 Sep 2026', '08 Sep 2026', '15 Sep 2026', 'Held'))
        ayesha, tanvir = values[header_index + 1], values[header_index + 2]
        self.assertEqual(ayesha[1:10], (2302001, 'Ayesha Rahman', 'a@example.com', 'P', 'A', 'P', 3, 2, 0.667))
        self.assertEqual(tanvir[4:10], (None, 'P', 'A', 2, 1, 0.5))
        self.assertEqual(ws.freeze_panes, f'E{header_index + 2}')

    def test_spreadsheets_never_turn_a_name_into_a_formula(self):
        from openpyxl import load_workbook

        self.ayesha.user.first_name = '=HYPERLINK("http://bad.example","x")'
        self.ayesha.user.save()
        rows = list(csv.reader(io.StringIO(self.export('format=csv').content.decode())))
        self.assertTrue(rows[1][1].startswith("'="), rows[1][1])
        ws = load_workbook(io.BytesIO(self.export('format=xlsx').content)).active
        names = [cell.value for cell in ws['C'] if cell.value and 'HYPERLINK' in str(cell.value)]
        self.assertEqual([cell.data_type for cell in ws['C'] if cell.value in names], ['s'])
        self.assertTrue(names[0].startswith("'="))

    def test_pdf_header_and_dhaka_time(self):
        with mock.patch('django.utils.timezone.now', return_value=FIXED_NOW):
            content = self.export('format=pdf').content
        text = pdf_text(content)
        self.assertTrue(content.startswith(b'%PDF'))
        self.assertIn('(CSE301 \\227 Software Engineering)', text)
        self.assertIn('Semester: Level 3 \\267 Term I \\267 2025-26', text)
        self.assertIn('Department: Computer Science and Engineering \\(CSE\\)', text)
        self.assertIn('Teacher: Karim Uddin', text)
        self.assertIn('06 Oct 2026, 02:30 AM \\(Dhaka time\\)', text)
        self.assertIn('(66.7%)', text)
        self.assertEqual(pdf_pages(content), 1)

    def test_pdf_splits_many_dates_so_id_name_and_percent_show_on_every_page(self):
        per_page = dates_per_page()
        start = date(2026, 1, 1)
        days = [start + timedelta(days=i) for i in range(per_page * 2 + 3)]  # three parts
        for day in days:
            log(self.ci, self.ayesha, day, 'PRESENT')
        content = self.export('format=pdf').content
        self.assertEqual(pdf_pages(content), 3)
        pages = pdf_text(content).split('(Student ID) Tj')[1:]
        self.assertEqual(len(pages), 3)
        for page in pages:
            self.assertIn('(Ayesha Rahman)', page)
            self.assertIn('(Tanvir Hasan)', page)
            self.assertIn('(%)', page)
        total = len(days) + 3
        self.assertIn(f'Classes 1\\226{per_page} of {total}', pdf_text(content))

    def test_pdf_one_date_and_empty_course(self):
        res = self.export('format=pdf&date=2026-09-08')
        self.assertEqual(res.status_code, 200)
        self.assertIn('(Present)', pdf_text(res.content))
        StudentClassroom.objects.all().delete()
        AttendanceLog.objects.all().delete()
        res = self.export('format=pdf')
        self.assertIn('No students in this course yet.', pdf_text(res.content))

    def test_docx(self):
        from docx import Document

        with mock.patch('django.utils.timezone.now', return_value=FIXED_NOW):
            document = Document(io.BytesIO(self.export('format=docx').content))
        paragraphs = [p.text for p in document.paragraphs]
        self.assertIn('CSE301 — Software Engineering', paragraphs)
        self.assertIn('Semester: Level 3 · Term I · 2025-26', paragraphs)
        self.assertIn('Department: Computer Science and Engineering (CSE)', paragraphs)
        self.assertIn('Teacher: Karim Uddin', paragraphs)
        self.assertIn('Generated: 06 Oct 2026, 02:30 AM (Dhaka time)', paragraphs)
        self.assertIn('Karim Uddin', document.sections[0].header.paragraphs[0].text)
        [table] = document.tables
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        self.assertEqual(rows[0], ['#', 'Student ID', 'Name', '01 Sep', '08 Sep', '15 Sep', 'Held', 'Present', '%'])
        self.assertEqual(rows[1], ['1', '2302001', 'Ayesha Rahman', 'P', 'A', 'P', '3', '2', '66.7%'])
        self.assertEqual(rows[2], ['2', '2302002', 'Tanvir Hasan', '', 'P', 'A', '2', '1', '50.0%'])

    def test_docx_splits_many_dates_into_tables(self):
        from docx import Document

        for i in range(dates_per_page() + 1):
            log(self.ci, self.ayesha, date(2026, 1, 1) + timedelta(days=i), 'PRESENT')
        document = Document(io.BytesIO(self.export('format=docx').content))
        self.assertEqual(len(document.tables), 2)
        for table in document.tables:
            header = [cell.text for cell in table.rows[0].cells]
            self.assertEqual(header[:3], ['#', 'Student ID', 'Name'])
            self.assertEqual(header[-1], '%')

    def test_docx_one_date(self):
        from docx import Document

        document = Document(io.BytesIO(self.export('format=docx&date=2026-09-08').content))
        rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
        self.assertEqual(rows[0], ['#', 'Student ID', 'Name', 'Email', 'Status'])
        self.assertEqual(rows[1][4], 'Absent')
        self.assertEqual(rows[2][4], 'Present')
